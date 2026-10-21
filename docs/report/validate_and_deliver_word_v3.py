"""Validate the rendered, visually reviewed v3 and deliver without overwriting."""
from pathlib import Path
from hashlib import sha256
from zipfile import ZipFile
import json
import re
import fitz
from docx import Document
from docx.oxml.ns import qn

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[1]
WORK=HERE/'word-v3'
SOURCE=ROOT/'output/20261021-阅记-项目报告-v2.docx'
RAW=ROOT/'raw/期末项目与报告要求.docx'
EDIT=WORK/'report.docx'
FINAL=ROOT/'output/20261021-阅记-项目报告-v3.docx'


def sha(path):
    return sha256(path.read_bytes()).hexdigest()


def main():
    baseline=json.loads((WORK/'content-checks.json').read_text(encoding='utf-8'))
    assert sha(SOURCE)==baseline['baseline_sha256']['source_v2']
    assert sha(RAW)==baseline['baseline_sha256']['raw']
    with ZipFile(EDIT) as z:
        assert z.testzip() is None
    old,new=Document(SOURCE),Document(EDIT)
    # Ignore recalculated TOC text only; compare all substantive body text.
    def body(d):
        return [p.text for p in d.paragraphs if p.text.strip() and not p.style.name.lower().startswith('toc ')]
    assert [t.replace('v2','v3') for t in body(old)]==body(new)
    def table_data(t):
        return [[c.text.replace('v2','v3') for c in r.cells] for r in t.rows]
    expected=[t for t in old.tables if not ('Library ViewModels' in '\n'.join(c.text for r in t.rows for c in r.cells) or '1:N' in '\n'.join(c.text for r in t.rows for c in r.cells))]
    assert [table_data(t) for t in expected]==[table_data(t) for t in new.tables]
    assert len(new.tables)==5 and len(new.inline_shapes)==21
    captions=[p.text for p in new.paragraphs if re.match(r'^图 \d+\s',p.text)]
    assert len(captions)==21
    old_media={sha256(p.blob).hexdigest() for p in old.part.package.parts if p.partname.startswith('/word/media/')}
    new_media={sha256(p.blob).hexdigest() for p in new.part.package.parts if p.partname.startswith('/word/media/')}
    assert old_media<=new_media
    svg_parts=[p for p in new.part.package.parts if p.content_type=='image/svg+xml']
    assert len(svg_parts)==2
    assert len(new.element.xpath('.//a:blip/a:extLst/a:ext'))==2
    assert not new.settings.element.xpath('.//w:documentProtection')
    assert not any('待确认' in t for t in body(new))
    render=json.loads((WORK/'native-render.json').read_text(encoding='utf-8'))
    assert render['pages']==22 and render['tables_of_contents']==1
    page_evidence=[]
    overflow=[]
    with fitz.open(WORK/'word-render.pdf') as pdf:
        assert len(pdf)==22
        for number,page in enumerate(pdf,1):
            assert (WORK/f'pages/page-{number:02}.png').is_file()
            text=page.get_text()
            assert text.strip()
            for block in page.get_text('dict')['blocks']:
                for line in block.get('lines',[]):
                    for span in line['spans']:
                        box=fitz.Rect(span['bbox'])
                        if box.x0 < -1 or box.y0 < -1 or box.x1>page.rect.width+1 or box.y1>page.rect.height+1:
                            overflow.append({'page':number,'text':span['text'],'bbox':list(box)})
            page_evidence.append({'page':number,'text_length':len(text),'raster_images':len(page.get_images())})
        assert not overflow,overflow
        # Vector artwork stays vector in the actual Word PDF output.
        assert not pdf[3].get_images() and not pdf[4].get_images()
        assert '图 1' in pdf[3].get_text() and '图 2' in pdf[4].get_text()
    assert 'SVG' not in '\n'.join(body(new))
    # User explicitly requested v3; deliver only after the final PNG review.
    with FINAL.open('xb') as dst:
        dst.write(EDIT.read_bytes())
    assert sha(FINAL)==sha(EDIT)
    verified=Document(FINAL)
    assert len(verified.inline_shapes)==21 and len(verified.tables)==5
    meta={'delivered':str(FINAL),'sha256':sha(FINAL),'size_bytes':FINAL.stat().st_size,'pages':22,'all_body_text_preserved_except_version':True,'remaining_data_tables_preserved':True,'all_19_original_images_preserved':True,'diagram_vectors':len(svg_parts),'diagrams_no_longer_tables':True,'captions':21,'toc_and_pages_updated_by_native_word':True,'all_22_final_page_pngs_visually_reviewed':True,'no_outside_page_text':True,'raw_and_v2_hash_unchanged':True,'protection':False,'unconfirmed_fields':0,'page_evidence':page_evidence}
    (WORK/'validation.json').write_text(json.dumps(meta,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:v for k,v in meta.items() if k!='page_evidence'},ensure_ascii=False,indent=2))


if __name__=='__main__':
    main()
