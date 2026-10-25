from pathlib import Path
from collections import Counter
import json,hashlib,zipfile,subprocess,shutil
import fitz
from docx import Document
from docx.oxml.ns import qn
from PIL import Image,ImageDraw

ROOT=Path(__file__).resolve().parents[3];WORK=Path(__file__).resolve().parent
def digest(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def norm(s):return ''.join(s.split())
facts=json.loads((WORK/'facts.json').read_text())
content=json.loads((WORK/'document-content.json').read_text())
expected=json.loads((WORK/'toc-pages.json').read_text())
before=json.loads((WORK/'protected-before.json').read_text())
assert all((ROOT/k).is_file() and digest(ROOT/k)==v for k,v in before.items())
head=subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip()
remote=subprocess.check_output(['git','rev-parse','origin/main'],cwd=ROOT,text=True).strip()
assert head==remote==facts['commit']
source=json.loads((ROOT/'docs/evidence/integrated-audit-20261024/source-hashes.json').read_text(encoding='utf-8-sig'))
assert all(digest(ROOT/r['path'])==r['sha256'] for r in source)
name='20261025-阅记-项目报告-v2'
word=WORK/(name+'.docx');pdf=WORK/(name+'.pdf')
d=Document(word);book=fitz.open(pdf)
assert len(book)==29,len(book)
assert len(d.inline_shapes)==11 and len(d.tables)==15
text='\n'.join(p.get_text() for p in book);body='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for r in t.rows for c in r.cells)
for field in ['周宇斌','20246085','软件2401','2026 年 10 月 25 日','ba08e19','353','351','137','41','344074','208','AppMetadata']:
 assert norm(field) in norm(body),(field,'missing')
assert '2026 年 10 月 21 日' not in body
assert '不提供电子书阅读器' not in body
assert '待确认' not in body
assert '\ufffd' not in text
for title,level,anchor in content['headings']:
 rows=[r for r in book.get_toc() if norm(r[1])==norm(title)]
 assert len(rows)==1 and rows[0][2]==expected[title],(title,rows,expected[title])
 assert any(e.attrib.get(qn('w:name'))==anchor for e in d.element.iter(qn('w:bookmarkStart')))
for n in range(1,12):assert ('图'+str(n)) in norm(text),('missing figure',n)
# Ensure no table caption/code/image creates a nearly empty spill page.
page_checks=[];geometry=[]
pages_dir=WORK/'pages-final';pages_dir.mkdir(exist_ok=True)
for i,page in enumerate(book):
 txt=page.get_text();assert len(txt)>85,(i+1,'nearly empty page')
 for block in page.get_text('dict')['blocks']:
  if block['type']!=0:continue
  for line in block['lines']:
   x0,y0,x1,y1=line['bbox']
   if x0<35 or x1>page.rect.width-35 or y0<12 or y1>page.rect.height-12:geometry.append({'page':i+1,'bbox':line['bbox'],'text':''.join(s['text'] for s in line['spans'])})
 pix=page.get_pixmap(matrix=fitz.Matrix(1.2,1.2),alpha=False);pix.save(pages_dir/f'page-{i+1:02}.png')
 page_checks.append({'page':i+1,'characters':len(txt),'links':len(page.get_links()),'image_instances':len(page.get_image_info())})
assert not geometry,geometry
assert len(book[1].get_links())>=19,'TOC not linked'
assert len(book.get_toc())==19
(WORK/'report-text.txt').write_text('\n'.join(line.rstrip() for line in text.splitlines())+'\n')
with zipfile.ZipFile(word) as z:
 assert len([n for n in z.namelist() if n.endswith('.svg')])==2
 assert 'TOC' in z.read('word/document.xml').decode()
 for n in z.namelist():
  if n.endswith('.xml'):
   from lxml import etree
   etree.fromstring(z.read(n))
package=Path(facts['package']['Path'])
with zipfile.ZipFile(package) as z:
 assert z.testzip() is None
 entries=(ROOT/'Code/delivery-files.txt').read_text().splitlines();assert z.namelist()==entries
 assert len(entries)==208
 for entry in entries:assert hashlib.sha256(z.read(entry)).hexdigest()==digest(ROOT/'Code'/entry)
 assert not any(any(p in ['bin','obj','.git','TestResults'] for p in Path(n).parts) for n in entries)
assert package.stat().st_size==344074 and package.stat().st_size<10_000_000
thumbs=[]
for i in range(len(book)):
 im=Image.open(pages_dir/f'page-{i+1:02}.png');im.thumbnail((290,410));tile=Image.new('RGB',(310,440),'#ddd');tile.paste(im,((310-im.width)//2,20));ImageDraw.Draw(tile).text((8,3),str(i+1),fill='black');thumbs.append(tile)
for start in range(0,len(thumbs),12):
 sheet=Image.new('RGB',(1240,1320),'white')
 for j,tile in enumerate(thumbs[start:start+12]):sheet.paste(tile,((j%4)*310,(j//4)*440))
 sheet.save(WORK/f'final-contact-{start//12+1}.jpg')
out=ROOT/'output/v2'
published=[]
for p in [word,pdf]:
 target=out/p.name
 if target.exists():assert digest(target)==digest(p),'Refusing to overwrite a different delivery'
 else:
  with target.open('xb') as f:f.write(p.read_bytes())
 assert digest(target)==digest(p)
 published.append({'path':str(target.relative_to(ROOT)),'bytes':target.stat().st_size,'sha256':digest(target)})
published.append({'path':str(package.relative_to(ROOT)),'bytes':package.stat().st_size,'sha256':digest(package)})
result={'passed':True,'commit':head,'pages':len(book),'headings':19,'native_tables':15,'figures':11,'code_excerpts':7,'svg_diagrams':2,'toc_pages_verified':expected,'toc_hyperlinks':len(book[1].get_links()),'all_pages_rendered':True,'page_checks':page_checks,'out_of_page_text':geometry,'protected_files_unchanged':len(before),'source_evidence_files_matched':len(source),'package_content_matched':208,'new_application_test_run':False,'visual_review':'See all 29 pages and final-contact-1/2/3.jpg; detailed diagrams/code/tables also inspected.','deliveries':published}
(WORK/'validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
print(json.dumps({k:v for k,v in result.items() if k not in ['page_checks','toc_pages_verified']},ensure_ascii=False,indent=2))
