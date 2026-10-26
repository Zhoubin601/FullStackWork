from pathlib import Path
from collections import Counter
import json,hashlib,zipfile,re,shutil,sys
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader
import pdfplumber
from PIL import Image,ImageDraw

W=Path(__file__).resolve().parent;ROOT=W.parents[2];OUT=ROOT/'output/提交材料'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def norm(s):return ''.join(s.split())
inputs=json.loads((W/'inputs.json').read_text(encoding='utf-8'))
assert all(sha(ROOT/k)==v for k,v in inputs['protected'].items())
word=W/'阅记-项目报告.docx';pdf=W/'阅记-项目报告.pdf';archive=W/'阅记-源码.zip'
d=Document(word);reader=PdfReader(pdf)
assert len(d.tables)==13 and len(d.inline_shapes)==10
before_doc=Document(W/'before/阅记-项目报告.docx')
removed_blip=before_doc.inline_shapes[9]._inline.xpath('.//a:blip')[0]
removed_blob=before_doc.part.related_parts[removed_blip.get(qn('r:embed'))].blob
with zipfile.ZipFile(word) as current,zipfile.ZipFile(W/'before/阅记-项目报告.docx') as previous:
    actual=Counter(hashlib.sha256(current.read(n)).hexdigest() for n in current.namelist() if n.startswith('word/media/') and n.endswith('.png'))
    expected=Counter(hashlib.sha256(previous.read(n)).hexdigest() for n in previous.namelist() if n.startswith('word/media/') and n.endswith('.png'))
    expected.subtract([hashlib.sha256(removed_blob).hexdigest()])
    assert actual==+expected,'Remaining raster screenshots must be unchanged; removed local-path screenshot must not remain embedded'
doc_text='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for row in t.rows for c in row.cells)
pdf_text='\n'.join(p.extract_text() for p in reader.pages)
for text in [doc_text,pdf_text]:
    for unwanted in ['docs/','docs\\','raw/','output/','evidence/','integrated-audit','full-verified','backup-loop','real-whole','probe.json','actions.jsonl','证据','本轮','本次','底稿','提交整理','编写','用户保存','先失败后通过','五次真实请求','723','40624','Codex','ChatGPT','图 11','待确认']:
        assert unwanted not in text,unwanted
    assert not re.search('(?i)(^|[^a-z0-9])v[0-9]+',text)
    for required in ['137','41','353','351','SetParameter','ServiceLocator','ViewLocator','AI','10张业务表']:
        assert required in norm(text),required
    assert chr(0xfffd) not in text
for n in range(1,11):assert ('图'+str(n)) in norm(pdf_text)
native=json.loads((W/'native-render.json').read_text(encoding='utf-8-sig'))
assert len(reader.pages)==native['pages'] and len(native['headings'])==19
tocrows=[row for row in native['tocText'][0].split(chr(13)) if row.strip()]
for heading,row in zip(native['headings'],tocrows):
    title,page=row.rsplit(chr(9),1)
    assert norm(title)==norm(heading['title']) and int(page)==heading['page']
    assert norm(title) in norm(reader.pages[int(page)-1].extract_text())
def flatten(rows):
    for r in rows:
        if isinstance(r,list):yield from flatten(r)
        else:yield r
bookmarks=list(flatten(reader.outline));assert len(bookmarks)==19
for h,b in zip(native['headings'],bookmarks):
    assert norm(h['title'])==norm(b.title) and reader.get_destination_page_number(b)+1==h['page']
links=[a.get_object() for a in reader.pages[1].get('/Annots',[]) if a.get_object().get('/Subtype')=='/Link']
assert len(links)==19
ids={page.indirect_reference.idnum:i for i,page in enumerate(reader.pages,1)}
for h,link in zip(native['headings'],links):
    dest=link.get('/Dest') or link['/A']['/D'];assert ids[dest[0].idnum]==h['page']
checks=[]
with pdfplumber.open(pdf) as book:
    for n,page in enumerate(book.pages,1):
        text=page.extract_text() or '';assert len(text)>85,n
        assert not any(c['x0']<30 or c['x1']>page.width-30 or c['top']<10 or c['bottom']>page.height-10 for c in page.chars),n
        checks.append({'page':n,'characters':len(text),'images':len(page.images)})
    toc=book.pages[1]
    numbers=toc.filter(lambda x:x['object_type']=='char' and x['text'].isdigit() and x['x0']>490 and 90<x['top']<650).extract_words()
    assert len(numbers)==19 and max(x['x1'] for x in numbers)-min(x['x1'] for x in numbers)<1
fonts={}
for page in reader.pages:
    for ref in page['/Resources'].get('/Font',{}).values():
        f=ref.get_object();descendants=[v.get_object() for v in f.get('/DescendantFonts',[])] or [f]
        descriptors=[v.get('/FontDescriptor') for v in descendants]
        fonts[str(f.get('/BaseFont',''))]=all(x and any(k in x.get_object() for k in ['/FontFile','/FontFile2','/FontFile3']) for x in descriptors)
assert all(fonts.values()),fonts
with zipfile.ZipFile(archive) as z,zipfile.ZipFile(W/'before/阅记-源码.zip') as old:
    manifest=(ROOT/'Code/delivery-files.txt').read_text(encoding='utf-8').splitlines()
    assert z.namelist()==manifest==old.namelist() and len(manifest)==208 and z.testzip() is None
    changed=[]
    for name in manifest:
        assert hashlib.sha256(z.read(name)).hexdigest()==sha(ROOT/'Code'/name)
        if old.read(name)!=z.read(name):changed.append(name)
        assert not set(Path(name).parts)&{'bin','obj','.git','TestResults'}
    assert set(changed)=={'README.md','CHANGELOG.md'}
    for name in ['README.md','CHANGELOG.md']:
        text=z.read(name).decode('utf-8')
        assert not any(w in text for w in ['docs/','raw/','本轮','本次','已留证','证据','按用户'])
assert archive.stat().st_size<10_000_000
images=W/'pages';assert len(list(images.glob('page-*.png')))==len(reader.pages)
for start in range(0,len(reader.pages),4):
    sheet=Image.new('RGB',(1200,1740),'#ddd');draw=ImageDraw.Draw(sheet)
    for j,n in enumerate(range(start+1,min(start+5,len(reader.pages)+1))):
        im=Image.open(images/f'page-{n:02}.png').convert('RGB');im.thumbnail((580,820));x=(j%2)*600+(600-im.width)//2;y=(j//2)*870+25;sheet.paste(im,(x,y));draw.text(((j%2)*600+10,(j//2)*870+7),str(n),fill='black')
    sheet.save(W/f'contact-{start//4+1}.jpg',quality=94)
(W/'report-text.txt').write_text(pdf_text,encoding='utf-8')
result={'passed':True,'pages':len(reader.pages),'tables':13,'figures':10,'toc_and_bookmarks':19,'no_external_process_references':True,'api_usage_and_cleanup_logs_removed':True,'embedded_fonts':fonts,'page_checks':checks,'runtime_and_tests_unchanged':True,'new_runtime_test_run':False,'package':{'members':208,'bytes':archive.stat().st_size,'sha256':sha(archive)},'protected_files_unchanged':True}
(W/'validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in result.items() if k not in ['embedded_fonts','page_checks']},ensure_ascii=False,indent=2))
if '--publish' in sys.argv:
    visual=json.loads((W/'visual-review.json').read_text(encoding='utf-8'))
    assert visual['passed'] and visual['pages']==len(reader.pages)
    assert visual['pdf_sha256']==sha(pdf) and visual['docx_sha256']==sha(word)
    assert all(sha(OUT/k)==v for k,v in inputs['deliveries'].items()),'User edited the current delivery during cleanup'
    deliveries=[]
    for file in [word,pdf,archive]:
        shutil.copy2(file,OUT/file.name);assert sha(file)==sha(OUT/file.name)
        deliveries.append({'path':str(OUT/file.name),'bytes':file.stat().st_size,'sha256':sha(file)})
    (W/'deliveries.json').write_text(json.dumps(deliveries,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(deliveries,ensure_ascii=False,indent=2))
