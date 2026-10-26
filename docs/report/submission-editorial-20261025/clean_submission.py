from pathlib import Path
from copy import deepcopy
import hashlib,json,shutil,subprocess,sys,zipfile
from docx import Document

ROOT=Path(__file__).resolve().parents[3];W=Path(__file__).resolve().parent
OUT=ROOT/'output/提交材料';BEFORE=W/'before'
BEFORE.mkdir(exist_ok=True)
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
if (W/'inputs.json').exists():raise RuntimeError('Inputs already captured; do not rerun over user changes')
original={}
for p in OUT.iterdir():
    if p.is_file():
        original[p.name]=sha(p);shutil.copy2(p,BEFORE/p.name)
for name in ['README.md','CHANGELOG.md']:shutil.copy2(ROOT/'Code'/name,BEFORE/name)
protected={p.relative_to(ROOT).as_posix():sha(p) for p in list((ROOT/'raw').rglob('*'))+list((ROOT/'output/v2').rglob('*')) if p.is_file()}
(W/'inputs.json').write_text(json.dumps({'deliveries':original,'protected':protected},ensure_ascii=False,indent=2),encoding='utf-8')

# Source package instructions must stand on their own alongside the report.
path=ROOT/'Code/README.md';text=path.read_text(encoding='utf-8')
paras=text.split('\n\n')
for i,p in enumerate(paras):
    if p.startswith('已修复缺失字体阻断排版'):
        paras[i]='Windows Release 自动化测试共353项：351通过、0失败、2项因符号链接权限跳过。测试覆盖阅读业务、数据存储、导航、原生阅读、AI核验和备份恢复；测试方法、结果及界面示例见配套项目报告。'
    elif p.startswith('## 本轮字体与 AI 调整'):
        paras[i]='## 字体与 AI 行为'
    elif p.startswith('Windows 实际窗口验收覆盖'):
        paras[i]='界面测试覆盖手动书籍、阅读记录、书摘收藏、阅读计划、混合统计、TXT/EPUB导入、分页、目录、原生选区、续读、完成与重置、AI原文核验和撤销，以及完整备份恢复。恢复前后10张业务表内容及12个受管文件摘要一致。'
    elif p.startswith('已实测Windows150%缩放'):
        paras[i]='已实测Windows 150%缩放；100%/200%与跨屏尚未验证。EPUB测试使用自建样本，未覆盖商业书籍的全部排版。复杂CSS、固定版式、DRM、非UTF-8 XML及SVG/公式/音视频暂不支持或明确降级，其他平台和图像格式的支持范围以实际测试为准。'
    elif p.startswith('2026-10-23官方配置依据'):
        paras[i]='服务地址为 https://api.deepseek.com/chat/completions ，默认deepseek-flash，可选deepseek-v4-pro；采用非流式JSON响应并显式关闭thinking。旧deepseek-chat/deepseek-reasoner配置仍可读取，发送前提示更新模型，不静默替换选择。实际模型可用性和费用以供应商说明为准。'
    elif p.startswith('`python pack.py'):
        paras[i]='`python pack.py <新ZIP路径>`按delivery-files.txt显式白名单打包，拒绝覆盖已有ZIP，并核对全部文件内容与10,000,000字节上限。压缩包只包含源码、资源、构建脚本和依赖说明，不加入bin/obj、运行数据库、凭据或测试数据。依赖及许可证见DEPENDENCIES.md和THIRD-PARTY-NOTICES.md。'
path.write_text('\n\n'.join(paras),encoding='utf-8',newline='')
path=ROOT/'Code/CHANGELOG.md';text=path.read_text(encoding='utf-8')
text=text.replace('2026-10-24 Windows Release 自动化353项：351通过、0失败、2项符号链接权限跳过；153项生产检查、14项原生排版实验、5次真实AI调用以及实际窗口完整备份恢复已留证。本次提交前另对独立解包源码重新还原、构建和测试，详细结果见配套报告。','Windows Release 自动化测试共353项：351通过、0失败、2项符号链接权限跳过。测试涵盖业务规则、存储、导航、原生阅读、AI核验及备份恢复，具体结果见配套项目报告。')
path.write_text(text,encoding='utf-8',newline='')

d=Document(BEFORE/'阅记-项目报告.docx');paras=list(d.paragraphs);tables=list(d.tables)
def change(i,text):
    p=paras[i]
    if p.runs:
        p.runs[0].text=text
        for r in p.runs[1:]:r.text=''
    else:p.add_run(text)
replacements={
67:'应用 UI 与 Library 包含 137 个类、41 个接口、134 个 C# 源文件和 15 个视图文件。按工程、命名空间及完整类型名合并 partial 声明，统计手写 class 和 interface，不计入测试工程、自动生成代码、record、struct 和 enum。',
68:'课程参考项目 DailyPoetryA 包含 63 个类、10 个接口和 9 个视图。按类与接口数量不少于参考项目 70%、视图数量不少于参考项目的要求，本项目达到相应规模。15 个视图文件包括窗口、初始化、容器和记录弹窗，其中业务页面为 11 个。',
118:'以下界面展示书籍管理、原生阅读、AI辅助和数据恢复功能。界面中的书籍与摘录均为自建测试样本，记录数量用于说明操作效果。',
128:'重置阅读进度后，系统回到首章偏移 0，清除完成标记，并保留历史记录和书摘。删除书籍后，若来源文件仍在收件区，系统不会通过扫描自动重建该书籍；用户可通过文件选择器主动重新导入。',
142:'选区书摘以完整原文为一条记录，AI提供整体解读。测试中，选定的9字原文完整保存，重复生成未增加重复书摘；从书摘回到原文后复制内容一致，撤销生成后恢复到操作前的书摘数量。',
143:'当前页和页范围可以生成多条候选，每条均经连续原文核验；单独执行“AI解读”只显示结果，不写入书摘。应用通过DeepSeek接口获得响应，测试工程通过模拟HTTP处理器覆盖错误、取消及迟到响应等情况。',
149:'关闭并重新打开带图 EPUB 后，系统可根据非空 NonTextBlockKey 恢复图片所在位置。图片锚点与文字锚点共同提供稳定定位，动态页号只用于当前布局，避免字体或字号变化后续读错位。',
165:'两项跳过测试涉及受管内容、目的地和来源/收件区符号链接，原因是Windows账号没有创建符号链接的权限，因此不计入通过数量。自动化测试、架构集成测试、排版实验及界面测试采用不同方法，结果分别统计。',
167:'功能测试在独立数据目录中进行，覆盖正常输入、非法输入、重复操作、异常响应及关闭重启等情况。下表列出各场景的输入、预期行为和实际结果。',
176:'性能测试运行于Windows、.NET 8.0.30和32逻辑处理器环境。使用生产ReaderViewModel与ReaderSurface进行Skia离屏绘制，视口为760×440 DIP、输出1140×660像素。测量区分冷导入、预热打开、翻页和重排，不将离屏采样等同于物理显示缩放测试。',
178:'冷导入与预热首绘分别反映解析导入和缓存预热后的打开成本，不合并为同一首屏耗时。P95描述本组有限样本的分布，结果会受设备、书籍、字体和缓存状态影响。',
183:'图10和图11展示完整备份恢复及恢复后的阅读状态。测试流程为导出备份、删除电子书及关联数据、恢复备份、关闭并重启应用，再检查内容和原文定位。',
186:'恢复前后对比10张业务表中的记录及12个受管文件的SHA256摘要，结果一致。恢复后，本机深色主题保持，受管电子书、图片和书摘位置仍可使用。',
190:'恢复并重启后，再次打开EPUB，正文、图片和书摘原文定位均可使用。结合选区复制、书摘去重和批次撤销测试，验证了界面、数据库和受管文件之间的协作。',
194:'源码包在独立目录中完成锁定依赖还原、Release构建和自动化测试。353项测试中351项通过、0项失败、2项因符号链接权限跳过。界面测试覆盖主要阅读流程，备份恢复后10张业务表与12个受管文件一致，支持系统在已测试环境中的可用性。',
}
for i,text in replacements.items():change(i,text)
for i in [32,123,169,172,174,179,180,181,*range(196,205)]:
    p=paras[i]._p;p.getparent().remove(p)
for i in [12,14]:tables[i]._tbl.getparent().remove(tables[i]._tbl)

# Keep the course's test section self-contained and remove external log references.
def cell(t,r,c,text):
    p=tables[t].cell(r,c).paragraphs[0]
    if p.runs:
        p.runs[0].text=text
        for run in p.runs[1:]:run.text=''
    else:p.add_run(text)
    for extra in tables[t].cell(r,c).paragraphs[1:]:extra._p.getparent().remove(extra._p)
cell(5,8,1,'业务规则、数据存储、异步竞态与真实绑定的自动化测试；结合实际窗口验证主要操作流程。')
cell(8,0,1,'环境与方法')
cell(8,4,0,'测试日期');cell(8,4,1,'2026年10月24日至25日；界面功能测试与源码包独立构建测试。')
cell(8,5,0,'构建与运行');cell(8,5,1,'独立解包后使用 --locked-mode 还原依赖，Release构建并运行自动化测试；启动方法见源码包README。')
cell(8,6,0,'测试边界')
cell(9,1,1,'共353项：351通过、0失败、2项因符号链接权限跳过。')
cell(9,2,0,'架构集成与性能测试');cell(9,2,1,'153项通过，覆盖15个视图模型和36个服务接口；反复切书后文档、布局和解码图片引用归零。')
cell(9,5,0,'AI功能与数据恢复');cell(9,5,1,'解读、书摘生成、去重和撤销验证通过；恢复前后10张业务表记录及12个受管文件摘要一致。')
cell(10,10,2,'符合；迁移后原有字段保持，重复实例写入被阻止。')
cell(13,0,1,'测量结果')
for row in tables[13].rows:
    for c in row.cells:
        for p in c.paragraphs:
            for run in p.runs:run.text=run.text.replace('同一生产探针','同一组测试').replace('同一探针内','同一组测试中').replace('不冒充物理键鼠输入延迟','不等同于物理键鼠输入延迟')

# Project summary, rather than a list of report-writing inputs.
new=deepcopy(paras[193]._p)
for child in list(new):
    if child.tag.endswith('}r'):new.remove(child)
from docx.text.paragraph import Paragraph
newp=Paragraph(new,paras[193]._parent)
newp.add_run('系统实现的重点是保持数据语义与界面状态一致：手动页码和电子书百分比分开计算，阅读器通过稳定正文锚点适应重排；统一导航在切换页面前结束活动会话，事务与操作日志共同处理数据库和文件恢复。AI仅处理用户确认的范围，原文核验和本地重试降低了错误摘录及重复请求的风险。')
paras[194]._p.addnext(new)
for p in d.paragraphs:
    if p.text=='最终结果汇总':
        for r in p.runs:r.text=r.text.replace('最终结果汇总','测试结果汇总')
    if p.text.startswith('图 3  最新概览') or p.text.startswith('图 4  最新书库'):
        for r in p.runs:r.text=r.text.replace('最新概览','阅读概览').replace('最新书库','书库管理')
    if p.text=='（5）真实 AI 与整段选区书摘':
        for r in p.runs:r.text=r.text.replace('真实 AI','AI 解读')

d.core_properties.comments='';d.core_properties.keywords='';d.core_properties.last_modified_by='';d.core_properties.author='';d.core_properties.identifier=''
d.save(W/'阅记-项目报告.docx')
package=json.loads(subprocess.check_output([sys.executable,str(ROOT/'Code/pack.py'),str(W/'阅记-源码.zip')],encoding='utf-8'))
with zipfile.ZipFile(BEFORE/'阅记-源码.zip') as old,zipfile.ZipFile(W/'阅记-源码.zip') as newzip:
    assert old.namelist()==newzip.namelist()
    changed=[name for name in old.namelist() if old.read(name)!=newzip.read(name)]
    assert set(changed)=={'README.md','CHANGELOG.md'},changed
(W/'package.json').write_text(json.dumps(package|{'changed_members':changed,'runtime_files_unchanged':True,'reuse_independent_test':'351 passed / 0 failed / 2 skipped; prior independent build in this chat'},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'paragraphs':len(d.paragraphs),'tables':len(d.tables),'figures':len(d.inline_shapes),'package':package},ensure_ascii=False))
