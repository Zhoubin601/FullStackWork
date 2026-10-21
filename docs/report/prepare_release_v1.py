"""Prepare targeted OfficeCLI edits from the user's saved Word, never rebuild it."""
from pathlib import Path
import hashlib, json, re, shutil, zipfile
from collections import Counter
from xml.etree import ElementTree as ET
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
WORK = ROOT / 'docs/report/release-v1'
EVIDENCE = ROOT / 'docs/evidence/release-v1-20261021'
SOURCE = WORK / 'source-user-v3.docx'
TARGET = WORK / 'report.docx'
PACKAGE = WORK / '20261021-阅记-源码-v1.zip'
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()

def counters(path):
    element = ET.parse(path).find('.//{*}Counters')
    assert element is not None
    return {k: int(element.get(k, '0')) for k in ['total','passed','failed','notExecuted']}

workspace = counters(ROOT/'docs/evidence/back-navigation-20261021/after-fix.trx')
clean = counters(EVIDENCE/'clean-v1.trx')
assert workspace['total'] == workspace['passed'] == clean['total'] == clean['passed'] == 94
assert clean['failed'] == clean['notExecuted'] == 0
probe = json.loads((EVIDENCE/'architecture/runtime-checks.json').read_text(encoding='utf-8-sig'))
assert probe['passed'] and probe['checks_count'] == 71
with zipfile.ZipFile(PACKAGE) as z:
    assert z.testzip() is None
    entries=(ROOT/'Code/delivery-files.txt').read_text(encoding='utf-8').splitlines()
    assert z.namelist() == entries
    assert all(z.read(n)==(ROOT/'Code'/n).read_bytes() for n in entries)
assert PACKAGE.stat().st_size < 10_000_000
meta={'file':PACKAGE.name,'bytes':PACKAGE.stat().st_size,'files':len(entries),'sha256':sha(PACKAGE)}

doc=Document(SOURCE)
commands=[]
edits=[]
def setp(index, text, expected):
    p=doc.paragraphs[index-1]
    if expected not in p.text:
        matches=[(i,p) for i,p in enumerate(doc.paragraphs,1) if expected in p.text]
        assert len(matches)==1, (index,expected,[i for i,_ in matches])
        index,p=matches[0]
    pid=p._p.get(qn('w14:paraId'))
    path=f'/body/p[@paraId={pid}]' if pid else f'/body/p[{index}]'
    commands.append({'command':'set','path':path,'props':{'text':text}})
    edits.append({'paragraph':index,'before':p.text,'after':text})

setp(51, doc.paragraphs[50].text + '左上角返回仅在当前模块有上一页时显示，模块首页隐藏入口；编辑页仍执行未保存确认与忙碌保护。', '菜单提供')
setp(79, '以下十图保留2026年10月20日Linux真实X11窗口的功能展示，显示的是主动载入的虚构演示资料；它们属于历史版本，首页仍显示返回按钮，不代表本版入口状态，也不替代Windows验收。随后图13～20保留Windows业务验收证据，图21为本版源码独立解包首次启动，图22～23补充返回入口修复。非全部内容能在同一窗口高度内展示时，页面提供滚动；截图不补画隐藏内容。', '以下十图')
setp(118, doc.paragraphs[117].text.replace('当前93项','当前94项') + '本版另修复栈底仍显示无效返回入口的问题：CanGoBack从主VM现有栈派生，集合变化同步按钮可见性和命令可执行状态；新增真实MainView绑定回归，保留菜单清栈与内容逐级返回。', '后续修复还覆盖')
setp(123, '保留历史平台证据：2026年10月20日在Debian 13 / Linux x64、Xvfb 21.1.16 / 1440×1000环境完成原68项测试与桌面流程，后续收藏竞态修复版完成93项。历史数字仅对应当时版本；本版返回入口修复后的自动化测试为94项，验证平台为Windows，不冒称已重新执行全部Linux流程，也不将多个平台的同一测试相加。', '保留历史平台证据')
setp(124, 'Windows当前工作区完整94项通过；本次v1源码包在新目录独立解包，执行锁定还原、Release构建和完整测试，同样94通过、0失败、0跳过。独立构建0错误、3条NETSDK1206提示，来自固定SQLite依赖旧Alpine RID。还原使用本机现有NuGet配置与缓存，未隔离全部缓存，不证明首次还原可离线完成。', 'Windows工作区')
setp(126, f'正式交付统一放在output/v1/。配套源码为{meta["file"]}，共{meta["files"]}个白名单文件、{meta["bytes"]}字节，低于10MB；无bin、obj、编译结果或运行数据库。SHA256为{meta["sha256"]}。本轮验证CRC正常、清单一致、所有文件与当前Code逐字节一致。', '正式配套源码')
setp(127, '本次v1为统一交付版本，包含最新返回入口修复、对应回归和同步后的README，不沿用旧版包的测试结果代替本轮独立验证。包内源码已独立完成还原、构建和94项测试；解包应用在另一个空数据目录实际启动，初始化版本1数据库并正常关闭。所有编译结果和运行数据仅留过程验证目录，未回填源码ZIP。', '与实际独立测试')
setp(128, '本次对当前Windows Release程序集重新运行独立架构探针，71项断言全部通过，覆盖14个VM单例、14个服务接口、13个反射视图、资源Current、真实容器、12个路由常量、参数、重复目标及返回栈。该71项与94项自动化测试分别统计，不合并为业务单元测试数量。', '另对当前Windows')
setp(129, '本版独立解包应用在隔离空库启动、初始化和正常关闭成功，启动日志未见绑定错误；返回修复另完成书库→详情→编辑→返回详情→返回书库的Windows实际点击。此前Windows完整业务闭环与Linux历史证据继续按原范围保留；本轮不声称重新遍历所有历史UI边界，未执行Alpine、大数据压力及全部CPU/系统组合。', 'Windows三个工作区')
setp(132, '下表按本次独立解包TRX中的测试类归组统计，每条数据行对应实际执行用例；总计94项，全部通过。架构组新增返回入口与真实MainView命令绑定回归，其他组保留原有覆盖。', '下表按TRX')
setp(134, '工作区原始结果为docs/evidence/back-navigation-20261021/after-fix.trx；本轮独立源码结果为docs/evidence/release-v1-20261021/clean-v1.trx，配有restore.log、build.log和test.log。71项独立架构结果为同目录architecture/runtime-checks.json。所有数量依据实际记录，不以截图估算，也不把单次耗时作为性能保证。', '对应原始证据')
setp(135, 'Windows历史业务验收完成120页书籍、非法页码校验、10～20页/25分钟记录、第15页书摘、计划与收藏，核对统计后进行JSON导出、级联删除、恢复和深色重启。删除后四表全0，恢复后各1且逐字段一致，收藏保留；自动备份记录替换前空库，完整性正常、无外键错误。', 'Windows真实窗口另行完成')
setp(138, '图17～20保留Windows实际业务验收截图，图21更新为本次v1源码包独立解包后的首次启动，图22～23为返回入口修复的实际窗口。94项通过的数量依据上节两份TRX及本轮test.log，71项依据独立探针；不使用旧截图冒充最新测试终端结果。', '以下均为')
setp(150, '图 21  v1源码独立解包后的首次启动：隔离空库初始化成功', '图 21')
setp(151, '前述业务截图来自明确标注的虚构测试书，不是实际用户阅读资料。历史Windows逐字段核对结果保留于原验收证据；本版首次启动与返回验证使用另外的隔离数据目录，测试集不混用。本报告未对实际界面补画控件。', '截图对应业务数据')
setp(154, '阅记完成十个业务页面及阅读管理闭环，保留参考项目真实IoC、单例定位器、三层导航、统一参数入口、主ViewModel返回栈和反射视图缓存。本版Windows工作区和v1源码独立解包均94项通过，当前程序集71项独立架构断言通过；源码包也已实际启动验证。新增返回入口状态修复完成真实点击验证，既有业务闭环和Linux结果保留其原验证范围。', '阅记完成十个')
setp(155, '本统一交付版v1已同步最新返回入口、94项测试、独立源码验证和包元数据，交付目录含可编辑Word、正式PDF和源码ZIP；旧交付文件按用户要求清理。Windows通过结论来自实际验证。本机构建仍有固定SQLite依赖的Alpine RID提示，未执行Alpine兼容、大数据压力或多设备同步测试。后续可继续改善输入可用性并测量更大数据集；电子书阅读、联网检索和云同步不在本版范围内。', '本修订版')

commands += [
 {'command':'set','path':'/body/tbl[4]/tr[2]','props':{'c1':'架构10','c2':'真实容器、单例、构造注入、三层导航、反射缓存、失败重试、A/B切换、4个忙态路由，以及返回入口与实际MainView绑定。'}},
 {'command':'set','path':'/body/tbl[5]/tr[3]','props':{'c2':'Windows工作区与本次v1源码独立解包均94项通过；原有业务闭环及新增多级返回已实际验证。'}},
 {'command':'set','path':'/body/tbl[5]/tr[5]','props':{'c2':f'统一v1：Word、PDF及源码ZIP；源码{meta["files"]}文件、{meta["bytes"]}字节，逐文件一致。'}},
]
with zipfile.ZipFile(SOURCE) as z:
    headers=sorted(n for n in z.namelist() if re.fullmatch(r'word/header\d+\.xml',n))
    for i,name in enumerate(headers,1):
        xml=ET.fromstring(z.read(name))
        ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        text=''.join(t.text or '' for t in xml.findall('.//w:t',ns))
        if '修订版 v3' in text:
            commands.append({'command':'set','path':f'/header[{i}]/p[1]','props':{'text':'阅记 · 阅读与书摘管理\t项目报告  /  统一交付版 v1'}})

shutil.copy2(SOURCE,TARGET)
(WORK/'edits.json').write_text(json.dumps(commands,ensure_ascii=False,indent=2),encoding='utf-8')
(WORK/'changed-paragraphs.json').write_text(json.dumps(edits,ensure_ascii=False,indent=2),encoding='utf-8')
snapshot={'source_sha256':sha(SOURCE),'raw_sha256':sha(ROOT/'raw/期末项目与报告要求.docx'),'package':meta,'workspace_tests':workspace,'clean_tests':clean,'architecture_checks':71,'old_output_files':{p.name:sha(p) for p in (ROOT/'output').iterdir() if p.is_file()}}
(WORK/'baseline.json').write_text(json.dumps(snapshot,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'edits':len(commands),'package':meta,'tests':clean,'probe_checks':71},ensure_ascii=False,indent=2))
