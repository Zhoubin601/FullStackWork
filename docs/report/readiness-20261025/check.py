from pathlib import Path, PurePosixPath
from hashlib import sha256
from zipfile import ZipFile
from xml.etree import ElementTree as ET
import json, math, re
import fitz
from docx import Document
from pypdf import PdfReader
from PIL import Image, ImageDraw

WORK = Path(__file__).resolve().parent
ROOT = WORK.parents[2]
OUT = ROOT / 'output/提交材料'
OLD = ROOT / 'docs/report/submission-20261025'
EDIT = ROOT / 'docs/report/submission-editorial-20261025'
def sha(p): return sha256(p.read_bytes()).hexdigest()
def norm(s): return ''.join(s.split())
def readj(p): return json.loads(p.read_text(encoding='utf-8-sig'))

initial = {str(p.relative_to(ROOT)): sha(p) for folder in [ROOT/'raw', OUT] for p in folder.rglob('*') if p.is_file()}
deliveries = readj(EDIT/'deliveries.json')
assert all(sha(Path(x['path'])) == x['sha256'] for x in deliveries)
protected = readj(EDIT/'inputs.json')['protected']
# Historical output versions were removed by the user-authorized 2026-10-26
# cleanup. Original raw inputs remain mandatory and must be unchanged.
assert all(sha(ROOT/p) == v for p,v in protected.items() if p.startswith('raw/'))

archive = OUT/'阅记-源码.zip'
manifest = (ROOT/'Code/delivery-files.txt').read_text(encoding='utf-8').splitlines()
with ZipFile(archive) as z, ZipFile(OLD/'阅记-源码.zip') as old:
    assert z.testzip() is None and z.namelist() == manifest == old.namelist()
    assert len(manifest) == len(set(manifest)) == 208
    changed = []
    for name in manifest:
        p = PurePosixPath(name)
        assert not p.is_absolute() and '..' not in p.parts
        assert not set(p.parts) & {'bin','obj','.git','TestResults','docs','raw','output'}
        assert p.suffix.lower() not in {'.dll','.exe','.pdb','.db','.sqlite','.sqlite3','.user','.suo'}
        assert z.read(name) == (ROOT/'Code'/name).read_bytes(), name
        if z.read(name) != old.read(name): changed.append(name)
    assert set(changed) == {'README.md','CHANGELOG.md'}
    # All non-generated source/resource/project files must actually be packaged.
    missing = []
    source_extensions = {'.cs','.axaml','.csproj','.resx','.png','.ico','.ttf','.otf','.json','.props','.targets','.sln','.manifest'}
    for p in (ROOT/'Code').rglob('*'):
        relative = p.relative_to(ROOT/'Code')
        if p.is_file() and not set(relative.parts) & {'bin','obj','.git','TestResults'} and p.suffix.lower() in source_extensions:
            if relative.as_posix() not in manifest: missing.append(relative.as_posix())
    assert not missing, missing
    # Check relative project references from each packaged csproj.
    for name in manifest:
        if name.endswith('.csproj'):
            project = ET.fromstring(z.read(name))
            for ref in project.iter('ProjectReference'):
                target = (ROOT/'Code'/name).parent/ref.attrib['Include'].replace('\\','/')
                assert target.resolve().relative_to((ROOT/'Code').resolve()).as_posix() in manifest
    readme = z.read('README.md').decode('utf-8')
    assert all(s in readme for s in ['dotnet restore','--locked-mode','dotnet build','dotnet test','dotnet run','8.0.424'])
assert archive.stat().st_size < 10_000_000

summaries = {}
for filename, source_dir in [('application-inventory.json','Code'), ('reference-inventory.json','daily-poetry-a')]:
    inventory = readj(OLD/filename)
    for source in inventory['SourceFiles']:
        assert sha(ROOT/source_dir/source['Path']).lower() == source['Sha256'].lower()
    current_sources = {p.relative_to(ROOT/source_dir).as_posix() for p in (ROOT/source_dir).rglob('*.cs') if not set(p.parts)&{'bin','obj','.git'} and 'UnitTest' not in str(p)}
    assert current_sources == {s['Path'] for s in inventory['SourceFiles']}
    summaries[source_dir] = inventory['Summary']
app, ref = summaries['Code'], summaries['daily-poetry-a']
assert app['Classes'] >= math.ceil(ref['Classes']*.7)
assert app['Interfaces'] >= math.ceil(ref['Interfaces']*.7)
assert app['ViewFiles'] >= ref['ViewFiles']
test_summary = readj(OLD/'test-summary.json')
assert sha(OLD/test_summary['trx_file']) == test_summary['trx_sha256']
trx = ET.parse(OLD/test_summary['trx_file'])
ns = {'t':'http://microsoft.com/schemas/VisualStudio/TeamTest/2010'}
actual = trx.find('.//t:Counters',ns).attrib
assert actual == test_summary['counters'] and actual['passed']=='351' and actual['failed']=='0'
# The extracted build directory is a disposable cache. The retained original
# independently tested ZIP was compared byte for byte above; it is sufficient
# to reproduce the same extraction without keeping generated binaries.

word = OUT/'阅记-项目报告.docx'
with ZipFile(word) as z:
    assert z.testzip() is None
    xml = z.read('word/document.xml').decode('utf-8')
    assert '<w:ins ' not in xml and '<w:del ' not in xml
    assert 'word/comments.xml' not in z.namelist()
doc = Document(word)
assert len(doc.tables)==13 and len(doc.inline_shapes)==10
body = '\n'.join(p.text for p in doc.paragraphs)
table_text = '\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
required = ['项目背景与意义','需求分析','功能需求','非功能需求','系统设计','架构设计','功能模块设计','数据/数据库设计','接口与类设计','技术选型说明','系统实现','关键模块的核心实现','核心功能与UI截图','难点和解决方案','系统测试','测试环境','测试用例','测试截图','总结与展望']
reader = PdfReader(OUT/'阅记-项目报告.pdf')
assert len(reader.pages)==28
pdf_text = '\n'.join(p.extract_text() or '' for p in reader.pages)
for text in [body+'\n'+table_text, pdf_text]:
    assert all(norm(s) in norm(text) for s in required)
    assert not any(s in text for s in ['待确认','TODO','Codex','ChatGPT','docs/','raw/','output/','本轮','底稿'])
    assert '\ufffd' not in text
# Word stores a TOC tab before the page number; PDF renders a dotted leader.
# TOC text, page numbers and destinations are checked separately below.
missing_paragraphs = [p.text for p in doc.paragraphs if len(norm(p.text))>8 and not re.search(r'\t\d+$',p.text) and norm(p.text) not in norm(pdf_text)]
missing_cells = list({c.text for t in doc.tables for row in t.rows for c in row.cells if len(norm(c.text))>8 and norm(c.text) not in norm(pdf_text)})
assert not missing_paragraphs and not missing_cells, (missing_paragraphs, missing_cells)
native = readj(EDIT/'native-render.json')
def flatten(items):
    for x in items:
        if isinstance(x,list): yield from flatten(x)
        else: yield x
bookmarks = list(flatten(reader.outline))
headings = native['headings']
assert len(bookmarks)==len(headings)==19
for b,h in zip(bookmarks,headings):
    assert norm(b.title)==norm(h['title']) and reader.get_destination_page_number(b)+1==h['page']
    assert norm(h['title']) in norm(reader.pages[h['page']-1].extract_text())
tocrows=[x for x in native['tocText'][0].split(chr(13)) if x.strip()]
for row,h in zip(tocrows,headings):
    title,page=row.rsplit(chr(9),1)
    assert norm(title)==norm(h['title']) and int(page)==h['page']
links = [a.get_object() for a in reader.pages[1].get('/Annots',[]) if a.get_object().get('/Subtype')=='/Link']
assert len(links)==19
ids={page.indirect_reference.idnum:n for n,page in enumerate(reader.pages,1)}
for link,h in zip(links,headings):
    dest=link.get('/Dest') or link['/A']['/D']
    assert ids[dest[0].idnum]==h['page']
fonts={}
for page in reader.pages:
    for fref in page['/Resources'].get('/Font',{}).values():
        f=fref.get_object()
        descendants=[v.get_object() for v in f.get('/DescendantFonts',[])] or [f]
        fonts[str(f.get('/BaseFont',''))]=all(v.get('/FontDescriptor') and any(k in v['/FontDescriptor'].get_object() for k in ['/FontFile','/FontFile2','/FontFile3']) for v in descendants)
assert all(fonts.values())
page_checks=[]
pdf=fitz.open(OUT/'阅记-项目报告.pdf')
for i,page in enumerate(pdf,1):
    assert len(page.get_text().strip())>60
    for word_box in page.get_text('words'):
        rect=fitz.Rect(word_box[:4])
        assert rect.x0>=0 and rect.y0>=0 and rect.x1<=page.rect.width+1 and rect.y1<=page.rect.height+1
    page_checks.append({'page':i,'characters':len(page.get_text()),'images':len(page.get_images())})
    pix=page.get_pixmap(matrix=fitz.Matrix(1.4,1.4),alpha=False)
    pix.save(WORK/f'page-{i:02}.png')
for start in range(0,len(pdf),4):
    sheet=Image.new('RGB',(1200,1740),'#ddd'); draw=ImageDraw.Draw(sheet)
    for j,n in enumerate(range(start+1,min(start+5,len(pdf)+1))):
        im=Image.open(WORK/f'page-{n:02}.png').convert('RGB');im.thumbnail((580,820))
        x=(j%2)*600+(600-im.width)//2;y=(j//2)*870+25
        sheet.paste(im,(x,y));draw.text(((j%2)*600+10,(j//2)*870+7),str(n),fill='black')
    sheet.save(WORK/f'contact-{start//4+1}.jpg',quality=93)
(WORK/'report-text.txt').write_text(pdf_text,encoding='utf-8')
assert all(sha(ROOT/k)==v for k,v in initial.items())
result={'passed':True,'current_deliveries_match_last_validated_hashes':True,'package_bytes':archive.stat().st_size,'package_members':len(manifest),'package_matches_Code':True,'package_changes_from_independently_tested_package':changed,'all_other_members_match_independently_tested_baseline_zip':True,'new_runtime_test_run':False,'prior_test_results':test_summary['results'],'prior_test_date':'2026-10-25','prior_test_trx_hash_verified':True,'inventories_source_hashes_verified':True,'application':app,'reference':ref,'minimum_classes':math.ceil(ref['Classes']*.7),'minimum_interfaces':math.ceil(ref['Interfaces']*.7),'pdf_pages':28,'toc_links_bookmarks':19,'tables':13,'images':10,'word_all_paragraphs_and_cells_present_in_pdf':True,'embedded_fonts':fonts,'page_checks':page_checks,'raw_and_deliveries_unchanged':True,'input_hashes':initial,'visual_review_pending':True}
(WORK/'checks.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in result.items() if k not in ['embedded_fonts','page_checks','input_hashes']},ensure_ascii=False,indent=2))
