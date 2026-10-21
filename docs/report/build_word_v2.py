"""Create a native editable Word report from the reviewed report-v2 source."""
from pathlib import Path
import re,json
from PIL import Image as PILImage
from docx import Document
from docx.shared import Mm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[1]
OUT=HERE/'word-v2'
DOC=OUT/'report.docx'
doc=Document(DOC)
for child in list(doc._element.body):
    if child.tag!=qn('w:sectPr'):doc._element.body.remove(child)
for element in list(doc.sections[0]._sectPr):
    if element.tag==qn('w:pgNumType'):doc.sections[0]._sectPr.remove(element)
for element in list(doc.settings.element):
    if element.tag==qn('w:updateFields'):doc.settings.element.remove(element)
for name in ['Normal','Title','Heading 1','Heading 2','Heading 3','Caption','Header','Footer']:
    if name not in doc.styles:
        style=doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
        if name!='Normal':style.base_style=doc.styles['Normal']
for level,name in enumerate(['Heading 1','Heading 2','Heading 3']):
    pr=doc.styles[name].element.get_or_add_pPr()
    for element in list(pr):
        if element.tag==qn('w:outlineLvl'):pr.remove(element)
    outline=OxmlElement('w:outlineLvl');outline.set(qn('w:val'),str(level));pr.append(outline)

def font(style,size,bold=False,color='000000',latin='Calibri',east='宋体'):
    style.font.name=latin;style.font.size=Pt(size);style.font.bold=bold;style.font.color.rgb=RGBColor.from_string(color)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),east)

font(doc.styles['Normal'],11)
pf=doc.styles['Normal'].paragraph_format
pf.line_spacing=1.3;pf.space_after=Pt(6);pf.first_line_indent=Pt(22)
for name,size in [('Title',24),('Heading 1',20),('Heading 2',14),('Heading 3',12)]:
    font(doc.styles[name],size,True,east='黑体')
    p=doc.styles[name].paragraph_format;p.first_line_indent=Pt(0);p.space_before=Pt(12);p.space_after=Pt(8);p.keep_with_next=True
font(doc.styles['Caption'],9,color='586A62')
doc.styles['Caption'].paragraph_format.first_line_indent=Pt(0)
doc.styles['Caption'].paragraph_format.line_spacing=1.1
doc.styles['Caption'].paragraph_format.space_after=Pt(9)
for name in ['Header','Footer']:
    font(doc.styles[name],8.5,color='586A62')
    doc.styles[name].paragraph_format.first_line_indent=Pt(0)

def setup_section(section):
    section.page_width=Mm(210);section.page_height=Mm(297)
    section.left_margin=Mm(23);section.right_margin=Mm(23);section.top_margin=Mm(24);section.bottom_margin=Mm(20)
    section.header_distance=Mm(12);section.footer_distance=Mm(10)

setup_section(doc.sections[0])
doc.sections[0].different_first_page_header_footer=True
def addp(text='',style=None):return doc.add_paragraph(text,style)
def coverp(text,size,bold=False,before=0,after=0):
    p=addp();p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Mm(before);p.paragraph_format.space_after=Mm(after)
    r=p.add_run(text);r.font.size=Pt(size);r.bold=bold
    if bold:r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'黑体')
    return p
coverp('全栈开发技术期末项目报告',24,True,18,20).style=doc.styles['Title']
coverp('阅记',34,True,7,4)
coverp('阅读与书摘管理应用',18,False,0,4)
coverp('修订版 v2  /  配套源码 v5',11,False,0,18)
for label,value in [('项目名称','阅记阅读与书摘管理应用'),('姓名','周宇斌'),('学号','20246085'),('班级','软件2401')]:coverp(f'{label}：{value}',12,False,0,5)
coverp('2026 年 10 月 21 日',12,False,13,10)
coverp('依据实际源码、Linux历史记录及Windows本机验收修订。当前代码93项自动化测试、71项独立架构断言通过；配套源码v5，应用与测试源码同已验收的v4，仅README更新。',10,False,0,0)

main=doc.add_section(WD_SECTION_START.NEW_PAGE);setup_section(main)
main.different_first_page_header_footer=False
main.header.is_linked_to_previous=False;main.footer.is_linked_to_previous=False
header=main.header.paragraphs[0];header.text='阅记 · 阅读与书摘管理\t项目报告  /  修订版 v2'
header.paragraph_format.tab_stops.add_tab_stop(Mm(164),WD_ALIGN_PARAGRAPH.RIGHT)
footer=main.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
def field(paragraph,instruction,cache=''):
    for tag,text in [('begin',None),('instr',instruction),('separate',None),('cache',cache),('end',None)]:
        r=OxmlElement('w:r')
        if tag=='instr':
            e=OxmlElement('w:instrText');e.set(qn('xml:space'),'preserve');e.text=text
        elif tag=='cache':e=OxmlElement('w:t');e.text=text
        else:e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),tag)
        r.append(e);paragraph._p.append(r)
field(footer,' PAGE ','1')
pg=OxmlElement('w:pgNumType');pg.set(qn('w:start'),'1');main._sectPr.insert_element_before(pg,'w:cols','w:formProt','w:vAlign','w:noEndnote','w:titlePg','w:textDirection','w:bidi','w:rtlGutter','w:docGrid','w:printerSettings','w:sectPrChange')
toc_title=addp('目录');toc_title.paragraph_format.first_line_indent=Pt(0);toc_title.runs[0].font.size=Pt(20);toc_title.runs[0].bold=True
tocp=addp();tocp.paragraph_format.first_line_indent=Pt(0)
field(tocp,' TOC '+chr(92)+'o "1-2" '+chr(92)+'h '+chr(92)+'z '+chr(92)+'u ','目录')
doc.add_page_break()
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');doc.settings.element.insert_element_before(update,'w:hdrShapeDefaults','w:footnotePr','w:endnotePr','w:compat','w:docVars','w:rsids','m:mathPr','w:attachedSchema','w:themeFontLang','w:clrSchemeMapping','w:doNotIncludeSubdocsInStats','w:doNotAutoCompressPictures','w:forceUpgrade')

def clean(s):
    s=re.sub(r'\\path\{([^{}]*)\}',lambda m:m.group(1),s)
    for a,b in [(r'\newline','\n'),(r'\_','_'),(r'\#','#'),(r'\%','%'),(r'\&','&')]:s=s.replace(a,b)
    if re.search(r'\\[a-zA-Z]+',s):raise ValueError('Unknown inline LaTeX '+s)
    return s.replace('~',' ')

def shade(cell,fill):
    shd=OxmlElement('w:shd');shd.set(qn('w:val'),'clear');shd.set(qn('w:fill'),fill);cell._tc.get_or_add_tcPr().insert_element_before(shd,'w:noWrap','w:tcMar','w:textDirection','w:tcFitText','w:vAlign','w:hideMark','w:headers','w:cellIns','w:cellDel','w:cellMerge','w:tcPrChange')

def format_table(table,widths,header=True):
    table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    for col,width in zip(table.columns,widths):col.width=Mm(width)
    seen=[]
    for i,row in enumerate(table.rows):
        cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
        if i==0 and header:
            repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
        for j,cell in enumerate(row.cells):
            if cell._tc in seen:continue
            seen.append(cell._tc)
            span=cell._tc.get_or_add_tcPr().gridSpan
            span=span.val if span is not None else 1
            cell.width=Mm(sum(widths[j:j+span]));cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade(cell,'173D36' if i==0 and header else ('F3F6F4' if i%2 else 'FFFFFF'))
            pr=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
            for side,value in [('top','80'),('left','90'),('bottom','80'),('right','90')]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),value);e.set(qn('w:type'),'dxa');m.append(e)
            pr.insert_element_before(m,'w:textDirection','w:tcFitText','w:vAlign','w:hideMark','w:headers','w:cellIns','w:cellDel','w:cellMerge','w:tcPrChange')
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1.12;p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
                if header and j==len(widths)-1 and len(widths)>2:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size=Pt(10);r.bold=i==0 and header
                    r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 and header else '172720')
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    table._tbl.tblPr.insert_element_before(borders,'w:shd','w:tblLayout','w:tblCellMar','w:tblLook','w:tblCaption','w:tblDescription','w:tblPrChange')

figure_count=0;image_count=0;source_paragraphs=[];code_blocks=[]
def caption(text):
    global figure_count
    figure_count+=1;p=addp(f'图 {figure_count}  {text}','Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next=False

def diagram(architecture,caption_text):
    if architecture:
        texts=['UI：App资源、ViewLocator、Views、平台提示与文件选择','↓','Library ViewModels：构造注入、公开命令、统一参数入口','↓','Library Services：14项接口职责、阅读规则、存储、统计、备份','↓','SQLite：四类业务表、外键、版本、串行写入与事务']
        t=doc.add_table(rows=len(texts),cols=1)
        for row,text in zip(t.rows,texts):row.cells[0].text=text
        format_table(t,[164],header=False)
    else:
        t=doc.add_table(rows=3,cols=3)
        t.rows[0].cells[0].merge(t.rows[0].cells[2]).text='Book\n书籍与总页数'
        for j,label in enumerate(['1:N ↙','1:N ↓','1:N ↘']):t.cell(1,j).text=label
        for j,label in enumerate(['ReadingRecord\n实际页段、日期、分钟','Quote\n正文、感想、收藏、页码','ReadingPlan\n开始日与目标日']):t.cell(2,j).text=label
        format_table(t,[164/3]*3,header=False)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
    caption(caption_text)

def table(block):
    widths=[float(x) for x in re.findall(r'p\{([\d.]+)mm\}',block.split('\n')[0])]
    widths=[x*164/sum(widths) for x in widths]
    body=block.split(r'\toprule',1)[1].split(r'\bottomrule',1)[0].replace(r'\midrule','')
    rows=[]
    for row in body.split(chr(92)*2):
        if row.strip():rows.append([clean(x.strip()) for x in row.strip().split(' & ')])
    assert all(len(r)==len(widths) for r in rows)
    t=doc.add_table(rows=len(rows),cols=len(widths))
    for row,values in zip(t.rows,rows):
        for cell,value in zip(row.cells,values):cell.text=value
    format_table(t,widths)
    p=addp();p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(2);p.paragraph_format.first_line_indent=Pt(0)

def code(text):
    code_blocks.append(text);p=addp();p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.left_indent=Mm(3);p.paragraph_format.line_spacing=Pt(10.5);p.paragraph_format.keep_together=False
    r=p.add_run(text);r.font.name='Consolas';r.font.size=Pt(8);r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')

source=(HERE/'report-v2.tex').read_text(encoding='utf-8');source=source[source.index(r'\section*{项目背景与意义}'):]
lines=source.splitlines();i=0;section=subsection=0;pending=[];last_break=True
def flush():
    global last_break
    if pending:
        value=clean(''.join(pending));source_paragraphs.append(value);addp(value);pending.clear();last_break=False
while i<len(lines):
    line=lines[i].strip()
    if not line:flush();i+=1;continue
    h=re.match(r'\\(section\*?|subsection)\{([^{}]+)\}',line)
    if h:
        flush();kind,title=h.groups()
        if kind.startswith('section'):
            if kind=='section':section+=1;subsection=0;title=f'{section}  {title}'
            p=addp(title,'Heading 1')
            if title.startswith('4  '):p.paragraph_format.page_break_before=True
        else:subsection+=1;addp(f'{section}.{subsection}  {title}','Heading 2')
        last_break=False;i+=1;continue
    if line in [r'\clearpage',r'\newpage']:
        flush()
        if not last_break:doc.add_page_break()
        last_break=True;i+=1;continue
    if line.startswith(r'\end{document}'):flush();break
    if line.startswith(r'\begin{longtable}'):
        flush();block=[line];i+=1
        while r'\end{longtable}' not in block[-1]:block.append(lines[i]);i+=1
        table('\n'.join(block));last_break=False;continue
    if line.startswith(r'\begin{figure}'):
        flush();block=[line];i+=1
        while r'\end{figure}' not in block[-1]:block.append(lines[i]);i+=1
        block='\n'.join(block);label=re.search(r'\\caption\{([^{}]*)\}',block).group(1)
        diagram('架构图' in block,label);last_break=False;continue
    s=re.fullmatch(r'\\screen\{([^{}]+)\}\{([^{}]+)\}',line)
    if s:
        flush();path,label=s.groups();path=ROOT/'docs/evidence'/path
        with PILImage.open(path) as im:iw,ih=im.size
        scale=min(160/iw,90/ih)
        p=addp();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1
        r=p.add_run();pic=r.add_picture(str(path),width=Mm(iw*scale),height=Mm(ih*scale))
        pic._inline.docPr.set('descr',label);caption(label);image_count+=1;last_break=False;i+=1;continue
    inp=re.fullmatch(r'\\VerbatimInput\[([^\]]+)\]\{([^{}]+)\}',line)
    if inp:
        flush();opts,path=inp.groups();a=int(re.search(r'firstline=(\d+)',opts).group(1));b=int(re.search(r'lastline=(\d+)',opts).group(1))
        code('\n'.join((HERE/path).read_text(encoding='utf-8').splitlines()[a-1:b]));last_break=False;i+=1;continue
    if line.startswith(r'\begin{Verbatim}'):
        flush();i+=1;block=[]
        while not lines[i].strip().startswith(r'\end{Verbatim}'):block.append(lines[i]);i+=1
        code('\n'.join(block));last_break=False;i+=1;continue
    if line.startswith(chr(92)):raise ValueError('Unknown block '+line)
    pending.append(line);i+=1
flush()
doc.core_properties.title='阅记阅读与书摘管理应用项目报告'
doc.core_properties.author='周宇斌'
doc.core_properties.subject='报告v2可编辑Word版，配套源码v5'
doc.save(DOC)
check=Document(DOC)
plain='\n'.join(p.text for p in check.paragraphs)
assert all(p in plain for p in source_paragraphs)
assert all(c in plain for c in code_blocks)
assert len(check.inline_shapes)==19 and figure_count==21
assert len([p for p in check.paragraphs if p.style.name in ['Heading 1','Heading 2']])==19
(OUT/'content-checks.json').write_text(json.dumps({'source':'report-v2.tex','native_body_paragraphs_verified':len(source_paragraphs),'native_code_blocks':len(code_blocks),'native_heading_count':19,'native_table_count':len(check.tables),'pictures':image_count,'figure_captions':figure_count,'protection':False,'all_source_paragraphs_preserved':True},ensure_ascii=False,indent=2),encoding='utf-8')
print('Created native editable report: 19 headings, 19 pictures, 21 captions')
