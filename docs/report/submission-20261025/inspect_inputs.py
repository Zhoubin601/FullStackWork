from pathlib import Path
import hashlib,json,zipfile
from docx import Document
from lxml import etree

ROOT=Path(__file__).resolve().parents[3]
WORK=Path(__file__).resolve().parent
source=ROOT/'output/v2/20261025-阅记-项目报告-v2.docx'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
protected=list((ROOT/'raw').rglob('*'))+list((ROOT/'output').rglob('*'))
state={p.relative_to(ROOT).as_posix():sha(p) for p in protected if p.is_file()}
(WORK/'input-hashes.json').write_text(json.dumps(state,ensure_ascii=False,indent=2),encoding='utf-8')
for name,path in [('guidance',ROOT/'raw/期末项目与报告要求.docx'),('current',source)]:
    doc=Document(path)
    data={'paragraphs':[{'i':i,'style':p.style.name,'text':p.text} for i,p in enumerate(doc.paragraphs)],'tables':[[[c.text for c in r.cells] for r in t.rows] for t in doc.tables]}
    (WORK/f'{name}-content.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
    if name=='guidance':
        (WORK/'guidance-text.txt').write_text('\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(' | '.join(c.text for c in r.cells) for t in doc.tables for r in t.rows),encoding='utf-8')
    else:
        with zipfile.ZipFile(path) as z:
            for member in ['word/document.xml','word/styles.xml','word/numbering.xml']:
                if member not in z.namelist():continue
                tree=etree.fromstring(z.read(member))
                (WORK/Path(member).name).write_bytes(etree.tostring(tree,encoding='utf-8',pretty_print=True))
        print(json.dumps({'source_sha256':sha(path),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'styles':[(s.style_id,s.name) for s in doc.styles if 'TOC' in s.name or 'toc' in s.style_id.lower() or '标题' in s.name],'paragraph_texts':[{'i':i,'text':p.text} for i,p in enumerate(doc.paragraphs)]},ensure_ascii=False,indent=2))
