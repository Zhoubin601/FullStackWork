from pathlib import Path
import json,xml.etree.ElementTree as ET,re,zipfile,os
from collections import Counter
from docx import Document

WORK=Path(__file__).resolve().parent
trx=ET.parse(WORK/'test-results/submission.trx').getroot()
ns={'t':'http://microsoft.com/schemas/VisualStudio/TeamTest/2010'}
counters=trx.find('.//t:Counters',ns).attrib
results=Counter(r.attrib['outcome'] for r in trx.findall('.//t:UnitTestResult',ns))
assert int(counters['failed'])==0,counters
assert results=={'Passed':351,'NotExecuted':2},results
meta={'counters':counters,'results':dict(results),'skipped':[r.attrib['testName'] for r in trx.findall('.//t:UnitTestResult',ns) if r.attrib['outcome']=='NotExecuted'],'times':trx.find('t:Times',ns).attrib,'independent_source':True,'runtime_logic_changed':False}
(WORK/'test-summary.json').write_text(json.dumps(meta,ensure_ascii=False,indent=2),encoding='utf-8')
path=WORK/'阅记-项目报告.docx';d=Document(path)
def set_text(p,text):
    if p.runs:
        p.runs[0].text=text
        for r in p.runs[1:]:r.text=''
    else:p.add_run(text)
for p in d.paragraphs:
    if p.text.startswith('TRX 的 UnitTestResult'):
        set_text(p,'10 月 24 日验收与 10 月 25 日独立源码回归的 TRX 均含 351 个 Passed、2 个 NotExecuted，0 失败。跳过用例涉及受管内容、目的地及来源/收件区符号链接，原因是当前 Windows 账号没有相应创建权限，不计为通过。生产探针、原生实验和真实 API 属于不同检查集，不与自动化数量相加。')
    elif p.text.startswith('最终 Windows 验收在已覆盖'):
        set_text(p,'10 月 25 日提交源码包独立还原与 Release 构建成功，353 项回归中 351 通过、0 失败、2 项权限跳过，与 10 月 24 日结果一致。此前 Windows 实际窗口、153 项生产探针、14 项原生实验、5 次真实 AI 调用及完整备份恢复提供补充证据。缺失字体与首屏重排两处缺陷已修复并有永久回归；结论限于已覆盖样本和设备。')
    elif p.text.startswith('编辑底稿：'):
        set_text(p,'架构参考：课程提供的 DailyPoetryA 源码。项目保留真实依赖注入容器、应用资源定位器、三层导航、统一参数入口和主视图模型返回栈，在此基础上扩展原生阅读与 AI 工作流。')
d.tables[9].cell(1,0).text='完整自动化回归'
d.tables[9].cell(1,1).text='10月25日独立源码：353项，351通过、0失败、2项权限跳过；与10月24日验收计数一致。两轮分别留存TRX。'
d.save(path)
temporary=WORK/'svg-clean.docx'
count=0
with zipfile.ZipFile(path) as src,zipfile.ZipFile(temporary,'w',zipfile.ZIP_DEFLATED) as out:
    for info in src.infolist():
        data=src.read(info.filename)
        if info.filename.endswith('.svg'):
            cleaned=re.sub(rb'<!DOCTYPE[^>]*>',b'',data,flags=re.S)
            if cleaned!=data:count+=1
            data=cleaned
        out.writestr(info,data)
os.replace(temporary,path)
print(json.dumps({'tests':dict(results),'svg_dtd_removed':count},ensure_ascii=False))
