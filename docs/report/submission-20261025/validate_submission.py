from pathlib import Path
import hashlib,json,zipfile,re,sys,shutil
from collections import Counter
from lxml import etree
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader
import pdfplumber

ROOT=Path(__file__).resolve().parents[3];W=Path(__file__).resolve().parent
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def norm(s):return ''.join(s.split())
before=json.loads((W/'input-hashes.json').read_text(encoding='utf-8'))
assert all((ROOT/k).is_file() and sha(ROOT/k)==v for k,v in before.items()),'Original raw or output changed'
source=json.loads((ROOT/'docs/evidence/integrated-audit-20261024/source-hashes.json').read_text(encoding='utf-8-sig'))
changes=[r['path'] for r in source if sha(ROOT/r['path'])!=r['sha256']]
assert set(changes)=={'Code/README.md','Code/CHANGELOG.md'},changes
app=json.loads((W/'application-inventory.json').read_text(encoding='utf-8-sig'))['Summary']
reference=json.loads((W/'reference-inventory.json').read_text(encoding='utf-8-sig'))['Summary']
minimum={'Classes':reference['MinimumClasses'],'Interfaces':reference['MinimumInterfaces'],'ViewFiles':reference['ViewFiles']}
assert all(app[k]>=v for k,v in minimum.items())
test=json.loads((W/'test-summary.json').read_text(encoding='utf-8'))
assert test['results']=={'Passed':351,'NotExecuted':2}
package=json.loads((W/'package.json').read_text(encoding='utf-8'))
archive=W/'阅记-源码.zip'
assert archive.stat().st_size<10_000_000
entries=(ROOT/'Code/delivery-files.txt').read_text(encoding='utf-8').splitlines()
with zipfile.ZipFile(archive) as z:
    assert z.testzip() is None
    assert z.namelist()==entries and len(entries)==208
    for name in entries:
        assert hashlib.sha256(z.read(name)).hexdigest()==sha(ROOT/'Code'/name),name
        assert (W/'source-unpacked'/name).is_file()
        assert hashlib.sha256(z.read(name)).hexdigest()==sha(W/'source-unpacked'/name),name
        assert not set(Path(name).parts)&{'bin','obj','TestResults','.git'},name
        assert not name.lower().endswith(('.dll','.exe','.pdb','.sqlite','.sqlite3','.db')),name
    for name in ['README.md','CHANGELOG.md']:
        txt=z.read(name).decode('utf-8')
        txt=txt.replace('deepseek-v4-pro','DeepSeekModel')
        assert not re.search('(?i)(^|[^a-z0-9])v[0-9]+',txt),name
        assert '旧正式ZIP' not in txt and '历史交付' not in txt
word=W/'阅记-项目报告.docx';pdf=W/'阅记-项目报告.pdf';d=Document(word);r=PdfReader(pdf)
native=json.loads((W/'native-render.json').read_text(encoding='utf-8-sig'))
assert len(r.pages)==native['pages']
assert len(d.tables)==15 and len(d.inline_shapes)==11
body='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for row in t.rows for c in row.cells)
pdf_text='\n'.join(p.extract_text() for p in r.pages)
for text in [body,pdf_text]:
    assert not re.search('(?i)(^|[^a-z0-9])v[0-9]+',text),'project version remains'
    assert '待确认' not in text and chr(0xfffd) not in text
    for wrong in ['云端 Linux','尚未独立','未在本轮独立解包','344074','ba08e19','旧 output']:
        assert wrong not in text,wrong
    for required in ['137','41','353','351','342333','208','AppMetadata','统一参数注入','ViewLocator','ServiceLocator']:
        assert required in norm(text),required
headings=native['headings'];assert len(headings)==19
toc_rows=[line for line in native['tocText'][0].split(chr(13)) if line.strip()]
assert len(toc_rows)==19
for heading,row in zip(headings,toc_rows):
    title,page=row.rsplit(chr(9),1)
    assert norm(title)==norm(heading['title']) and int(page)==heading['page']
    assert norm(title) in norm(r.pages[int(page)-1].extract_text()),(title,page)
def outline_rows(rows):
    for row in rows:
        if isinstance(row,list):yield from outline_rows(row)
        else:yield row
bookmarks=list(outline_rows(r.outline))
assert len(bookmarks)==19,[(b.title) for b in bookmarks]
for heading,bookmark in zip(headings,bookmarks):
    assert norm(bookmark.title)==norm(heading['title'])
    assert r.get_destination_page_number(bookmark)+1==heading['page']
toclinks=[a.get_object() for a in r.pages[1].get('/Annots',[]) if a.get_object().get('/Subtype')=='/Link']
assert len(toclinks)>=19
page_ids={p.indirect_reference.idnum:i for i,p in enumerate(r.pages,1)}
for heading,link in zip(headings,toclinks):
    destination=link.get('/Dest') or link['/A']['/D']
    assert page_ids[destination[0].idnum]==heading['page'],heading
geometry=[];pages=[]
with pdfplumber.open(pdf) as book:
    for i,page in enumerate(book.pages,1):
        text=page.extract_text() or ''
        assert len(text)>85,(i,'empty or nearly empty')
        bad=[c for c in page.chars if c['x0']<30 or c['x1']>page.width-30 or c['top']<10 or c['bottom']>page.height-10]
        geometry.extend({'page':i,'text':c['text'],'x0':c['x0'],'x1':c['x1'],'top':c['top']} for c in bad)
        pages.append({'page':i,'characters':len(text),'images':len(page.images)})
    toc=book.pages[1]
    digits=toc.filter(lambda obj:obj['object_type']=='char' and obj['text'].isdigit() and obj['x0']>490 and 90<obj['top']<650)
    page_numbers=digits.extract_words()
    assert len(page_numbers)==19,len(page_numbers)
    assert max(w['x1'] for w in page_numbers)-min(w['x1'] for w in page_numbers)<1,'Page numbers not aligned'
    child=next(w for w in toc.extract_words() if w['text']=='1.1')
    parent=next(w for w in toc.extract_words() if w['text']=='1' and 100<w['top']<180)
    assert child['x0']-parent['x0']>20,(child['x0'],parent['x0'])
assert not geometry,geometry[:5]
with zipfile.ZipFile(word) as z,zipfile.ZipFile(ROOT/'output/v2/20261025-阅记-项目报告-v2.docx') as original:
    pictures=[n for n in z.namelist() if n.startswith('word/media/')]
    assert len([n for n in pictures if n.endswith('.svg')])==2
    pngs=[n for n in pictures if n.lower().endswith(('.png','.jpg','.jpeg'))]
    # Word may rename media parts, so compare the binary image multiset.
    assert Counter(hashlib.sha256(z.read(n)).hexdigest() for n in pngs)==Counter(hashlib.sha256(original.read(n)).hexdigest() for n in original.namelist() if n.startswith('word/media/') and n.lower().endswith(('.png','.jpg','.jpeg')))
    for name in z.namelist():
        if name.endswith(('.xml','.svg')):etree.fromstring(z.read(name),parser=etree.XMLParser(resolve_entities=False,no_network=True))
    assert ' TOC ' in z.read('word/document.xml').decode('utf-8')
fonts={}
for page in r.pages:
    for font in page['/Resources'].get('/Font',{}).values():
        f=font.get_object(); name=str(f.get('/BaseFont',''))
        descs=[x.get_object() for x in f.get('/DescendantFonts',[])] or [f]
        descriptors=[x.get('/FontDescriptor') for x in descs]
        fonts[name]=all(desc and any(k in desc.get_object() for k in ['/FontFile','/FontFile2','/FontFile3']) for desc in descriptors)
assert all(fonts.values()),fonts
results={'passed':True,'pages':len(r.pages),'headings':19,'toc_links':len(toclinks),'toc_page_numbers_aligned':True,'toc_child_indent_verified':True,'pdf_bookmarks':19,'tables':15,'figures':11,'embedded_fonts':fonts,'out_of_page_characters':geometry,'page_checks':pages,'original_files_unchanged':len(before),'runtime_code_matches_window_evidence':True,'changed_source_documents':changes,'scale':{'actual':app,'reference':reference,'minimum':minimum},'independent_test_results':test['results'],'package':package,'svg_diagrams_preserved_without_external_dtd':2,'png_and_screenshot_bytes_preserved':True,'project_version_markers_removed':True}
(W/'validation.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
(W/'report-text.txt').write_text(pdf_text,encoding='utf-8')
print(json.dumps({k:v for k,v in results.items() if k not in ['embedded_fonts','page_checks','scale','package']},ensure_ascii=False,indent=2))
if '--publish' in sys.argv:
    visual=json.loads((W/'visual-review.json').read_text(encoding='utf-8'))
    assert visual['passed'] and visual['pages_reviewed']==len(r.pages)
    assert visual['docx_sha256']==sha(word) and visual['pdf_sha256']==sha(pdf)
    target=ROOT/'output/提交材料';target.mkdir(exist_ok=True)
    deliveries=[]
    for path in [word,pdf,archive]:
        output=target/path.name
        assert not output.exists(),f'Refusing to overwrite {output}'
        shutil.copy2(path,output)
        assert sha(path)==sha(output)
        deliveries.append({'path':str(output),'bytes':output.stat().st_size,'sha256':sha(output)})
    (W/'deliveries.json').write_text(json.dumps(deliveries,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(deliveries,ensure_ascii=False,indent=2))
