"""Retain the user's layout, refresh one startup figure, append two real return screenshots."""
from pathlib import Path
from copy import deepcopy
import json, re, zipfile
from docx import Document
from docx.shared import Mm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parents[2]
WORK=ROOT/'docs/report/release-v1'
PATH=WORK/'report.docx'
doc=Document(PATH)

# An empty paragraph carrying a page break can itself spill to a blank page.
test_heading=next(p for p in reversed(doc.paragraphs) if p.text.strip()=='4.3  测试截图')
previous_break=test_heading._p.getprevious()
if previous_break is not None and previous_break.xpath('.//w:br[@w:type="page"]'):
    test_heading.paragraph_format.page_break_before=True
    previous_break.getparent().remove(previous_break)

caption=next(p for p in doc.paragraphs if p.text.startswith('图 21 '))
image_p=caption._p.getprevious()
while not image_p.xpath('.//a:blip'):
    image_p=image_p.getprevious()
rid,_=doc.part.get_or_add_image(str(ROOT/'docs/evidence/release-v1-20261021/clean-v1-first-launch.png'))
blips=image_p.xpath('.//a:blip')
assert len(blips)==1
old_rid=blips[0].get(qn('r:embed'))
old_startup_part=doc.part.related_parts[old_rid].partname
blips[0].set(qn('r:embed'),rid)
for node in image_p.xpath('.//wp:docPr'):
    node.set('descr','v1源码独立解包后的Windows首次启动，隔离空库初始化成功')

summary=next(p for p in reversed(doc.paragraphs) if p.text.startswith('5  总结与展望'))
anchor=summary._p
previous=anchor.getprevious()
if previous is not None and previous.xpath('.//w:br[@w:type="page"]'):
    anchor=previous

def before(text='',style='Normal'):
    element=OxmlElement('w:p')
    anchor.addprevious(element)
    paragraph=Paragraph(element,doc._body)
    paragraph.style=style
    if text:paragraph.add_run(text)
    return paragraph

intro=before('本次返回入口验证使用独立虚构书籍：从书库进入详情及编辑后逐级返回，页面栈保留；回到模块首页后不再展示无效入口。')
intro.paragraph_format.page_break_before=True
intro.paragraph_format.keep_with_next=True
intro.paragraph_format.space_after=Pt(5)
for image_path,label in [
    ('nested-editor.png','图 22  Windows返回修复：编辑资料页有上一页时显示返回入口'),
    ('library-after-detail-back.png','图 23  逐级返回书库首页后隐藏返回入口，测试书籍保留'),
]:
    p=before()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    f=p.paragraph_format
    f.first_line_indent=Pt(0)
    f.space_before=Pt(0)
    f.space_after=Pt(2)
    f.line_spacing=1
    f.keep_with_next=True
    picture=p.add_run().add_picture(str(ROOT/'docs/evidence/back-navigation-20261021'/image_path),width=Mm(145))
    picture._inline.docPr.set('descr',label)
    c=before(label,caption.style)
    c.runs[0].font.size=Pt(9)
    c.alignment=WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.first_line_indent=Pt(0)
    c.paragraph_format.space_before=Pt(0)
    c.paragraph_format.space_after=Pt(5)
    c.paragraph_format.keep_with_next=False

doc.save(PATH)
check=Document(PATH)
assert len(check.tables)==5 and len(check.inline_shapes)==23
body='\n'.join(p.text for p in check.paragraphs)
assert '当前93项' not in body and '总计94项' in body and '源码-v5.zip' not in body
assert not re.search(r'待确认|TODO|lorem|\{\{',body,re.I)
with zipfile.ZipFile(PATH) as z:
    svgs=[n for n in z.namelist() if n.endswith('.svg')]
    assert len(svgs)==2
info={'native_tables':len(check.tables),'inline_shapes':len(check.inline_shapes),'svg_diagrams':len(svgs),'figure_captions':len([p for p in check.paragraphs if re.match(r'^图 \d+',p.text)]),'startup_figure_refreshed':True,'old_startup_part':str(old_startup_part),'new_real_return_figures':2,'source':'source-user-v3.docx','no_template_rebuild':True}
(WORK/'content-checks.json').write_text(json.dumps(info,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(info,ensure_ascii=False,indent=2))
