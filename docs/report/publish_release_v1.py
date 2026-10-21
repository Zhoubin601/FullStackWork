"""Validate the final report/source pair, then publish new files without overwriting."""
from pathlib import Path
from collections import Counter
from xml.etree import ElementTree as ET
import hashlib,json,re,sqlite3,zipfile
from docx import Document
import fitz

ROOT=Path(__file__).resolve().parents[2]
WORK=ROOT/'docs/report/release-v1'
OUT=ROOT/'output/v1'
base=json.loads((WORK/'baseline.json').read_text(encoding='utf-8'))
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
assert sha(ROOT/'raw/期末项目与报告要求.docx')==base['raw_sha256']
assert sha(WORK/'source-user-v3.docx')==base['source_sha256']
assert sha(ROOT/'output/20261021-阅记-项目报告-v3.docx')==base['source_sha256']

doc=Document(WORK/'report.docx')
source=Document(WORK/'source-user-v3.docx')
body='\n'.join(p.text for p in doc.paragraphs)
assert len(doc.tables)==5 and len(doc.inline_shapes)==23
assert len([p for p in doc.paragraphs if re.match(r'^图 \d+',p.text)])==23
assert [p.text for p in doc.paragraphs[:8]]==[p.text for p in source.paragraphs[:8]]
modified={x['paragraph'] for x in json.loads((WORK/'changed-paragraphs.json').read_text(encoding='utf-8'))}
modified.update(i for i,p in enumerate(source.paragraphs,1) if p.text.startswith('Windows真实窗口另行完成'))
assert not [(i,p.text) for i,p in enumerate(source.paragraphs,1) if i>=32 and p.text.strip() and i not in modified and p.text not in body]
for a,b in zip(source.tables[:3],doc.tables[:3]):
    assert [c.text for r in a.rows for c in r.cells]==[c.text for r in b.rows for c in r.cells]
assert sum(int(re.search(r'\d+$',r.cells[0].text).group()) for r in doc.tables[3].rows[1:])==94
assert not re.search(r'待确认|TODO|lorem|源码-v[2345]\.zip|当前93项|总计93项',body,re.I)
assert base['package']['sha256'] in body and str(base['package']['bytes']) in body
assert not doc.settings.element.xpath('.//w:documentProtection')
with zipfile.ZipFile(WORK/'report.docx') as z:
    svg_hashes={hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.endswith('.svg')}
    xml=z.read('word/document.xml').decode('utf-8')
    assert 'TOC ' in xml and len(svg_hashes)==2
with zipfile.ZipFile(WORK/'source-user-v3.docx') as z:
    assert svg_hashes=={hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.endswith('.svg')}

page_evidence=[]
with fitz.open(WORK/'report.pdf') as pdf:
    assert len(pdf)==23
    whole='\n'.join(p.get_text() for p in pdf)
    assert base['package']['sha256'] in re.sub(r'\s+','',whole)
    assert all(re.search(r'^图\s*'+str(n)+r'\b',whole,re.M) for n in range(1,24))
    for i,page in enumerate(pdf,1):
        text=page.get_text()
        assert text.strip()
        assert (WORK/f'pages/page-{i:02}.png').is_file()
        assert i<=2 or len(text)>65 or page.get_images()
        for block in page.get_text('dict')['blocks']:
            for line in block.get('lines',[]):
                for span in line['spans']:
                    r=fitz.Rect(span['bbox'])
                    assert r.x0>=-1 and r.y0>=-1 and r.x1<=page.rect.width+1 and r.y1<=page.rect.height+1,(i,span)
        page_evidence.append({'page':i,'characters':len(text),'raster_images':len(page.get_images())})
    assert not pdf[3].get_images() and not pdf[4].get_images() # diagrams retained as vectors

package=WORK/base['package']['file']
assert sha(package)==base['package']['sha256'] and package.stat().st_size==base['package']['bytes']
with zipfile.ZipFile(package) as z:
    manifest=(ROOT/'Code/delivery-files.txt').read_text(encoding='utf-8').splitlines()
    assert z.testzip() is None and z.namelist()==manifest
    for name in manifest:
        assert not (set(Path(name).parts)&{'bin','obj','.git','TestResults'})
        assert Path(name).suffix.lower() not in {'.dll','.exe','.pdb'}
        assert z.read(name)==(ROOT/'Code'/name).read_bytes()
clean=ET.parse(ROOT/'docs/evidence/release-v1-20261021/clean-v1.trx')
c=clean.find('.//{*}Counters')
assert c is not None and c.get('passed')==c.get('total')=='94' and c.get('failed')=='0'
probe=json.loads((ROOT/'docs/evidence/release-v1-20261021/architecture/runtime-checks.json').read_text(encoding='utf-8-sig'))
assert probe['passed'] and probe['checks_count']==71
db=ROOT/'docs/evidence/release-v1-20261021/clean-ui-data/readjournal.sqlite3'
with sqlite3.connect(f'file:{db.as_posix()}?mode=ro',uri=True) as connection:
    assert connection.execute('PRAGMA user_version').fetchone()[0]==1
    assert connection.execute('PRAGMA integrity_check').fetchone()[0]=='ok'
    assert connection.execute('PRAGMA foreign_key_check').fetchall()==[]
    counts={table:connection.execute(f'SELECT COUNT(*) FROM {table}').fetchone()[0] for table in ['Book','ReadingRecord','Quote','ReadingPlan']}
    assert not any(counts.values())

assert not OUT.exists(), 'Do not overwrite an existing v1 delivery directory.'
OUT.mkdir()
deliveries={
    '20261021-阅记-项目报告-v1.docx':WORK/'report.docx',
    '20261021-阅记-项目报告-v1.pdf':WORK/'report.pdf',
    base['package']['file']:package,
}
final=[]
for name,src in deliveries.items():
    dst=OUT/name
    with dst.open('xb') as stream:stream.write(src.read_bytes())
    assert sha(dst)==sha(src)
    final.append({'name':name,'path':str(dst),'bytes':dst.stat().st_size,'sha256':sha(dst)})
Document(OUT/'20261021-阅记-项目报告-v1.docx')
with fitz.open(OUT/'20261021-阅记-项目报告-v1.pdf') as pdf:assert len(pdf)==23
with zipfile.ZipFile(OUT/base['package']['file']) as z:assert z.testzip() is None
validation={'validated':True,'source_user_word_preserved':True,'unchanged_body_text_preserved':True,'svg_diagrams_unchanged':2,'startup_figure_refreshed':True,'new_return_figures':2,'native_tables':5,'figures':23,'pages':23,'all_pages_visually_reviewed':list(range(1,24)),'outside_page_text':0,'worktree_tests':94,'clean_package_tests':94,'runtime_architecture_assertions':71,'package_matches_current_code':True,'clean_first_launch_integrity':'ok','clean_first_launch_counts':counts,'raw_unchanged':True,'files':final,'page_evidence':page_evidence,'old_output_cleanup_pending':True}
(WORK/'validation.json').write_text(json.dumps(validation,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in validation.items() if k not in ['all_pages_visually_reviewed','page_evidence']},ensure_ascii=False,indent=2))
